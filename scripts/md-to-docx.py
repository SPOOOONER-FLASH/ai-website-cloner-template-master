"""Render a Markdown file to .docx for the owner, whose computer cannot open .md.

  python3 scripts/md-to-docx.py <in.md> <out.docx>

Handles what our docs use: #–#### headings, paragraphs, pipe tables, - / 1. lists, fenced code,
inline **bold** and `code`. Requires python-docx. The .md stays the source; regenerate, never hand-edit.
"""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

src, out = sys.argv[1], sys.argv[2]
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
for s in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
    doc.styles[s].element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

def inline(p, text):
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part: continue
        if part.startswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        else:
            p.add_run(part)

def cells(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]

lines = open(src, encoding="utf-8").read().splitlines()
i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith("```"):
        i += 1; buf = []
        while i < len(lines) and not lines[i].startswith("```"):
            buf.append(lines[i]); i += 1
        r = doc.add_paragraph().add_run("\n".join(buf)); r.font.name = "Consolas"; r.font.size = Pt(9)
    elif m := re.match(r"^(#{1,4})\s+(.*)", ln):
        inline(doc.add_heading(level=len(m.group(1))), m.group(2))
    elif ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
        head = cells(ln); rows = []
        i += 2
        while i < len(lines) and lines[i].startswith("|"):
            rows.append(cells(lines[i])); i += 1
        t = doc.add_table(rows=1, cols=len(head)); t.style = "Table Grid"
        for c, h in zip(t.rows[0].cells, head):
            c.text = ""; inline(c.paragraphs[0], f"**{h}**" if h else "")
        for r in rows:
            rc = t.add_row().cells
            for c, v in zip(rc, r + [""] * (len(head) - len(r))):
                c.text = ""; inline(c.paragraphs[0], v)
        doc.add_paragraph()
        continue
    elif m := re.match(r"^\s*[-*]\s+(.*)", ln):
        inline(doc.add_paragraph(style="List Bullet"), m.group(1))
    elif m := re.match(r"^\s*\d+\.\s+(.*)", ln):
        inline(doc.add_paragraph(style="List Number"), m.group(1))
    elif ln.startswith(">"):
        p = doc.add_paragraph(); inline(p, ln.lstrip("> ")); p.runs and setattr(p.runs[0], "italic", True)
    elif ln.strip():
        # join wrapped continuation lines into one paragraph
        buf = [ln.strip()]
        while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(#|\||```|\s*[-*]\s|\s*\d+\.\s|>)", lines[i + 1]):
            i += 1; buf.append(lines[i].strip())
        inline(doc.add_paragraph(), "".join(buf) if re.search(r"[一-鿿]", buf[0]) else " ".join(buf))
    i += 1
doc.save(out)
print("wrote", out)
