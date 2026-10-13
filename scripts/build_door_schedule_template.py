#!/usr/bin/env python3
"""Writes the blank door hardware schedule attached to the door-schedule article.

openpyxl is already the xlsx writer here (scripts/build_master_key_workbook.py), so this
follows that file's look: one header block for the things a schedule states once (project,
revision, finish, keying), then one row per door. The columns are the ones the article
walks through; the help row under them says what to write, so the sheet works without the
article open beside it.

Usage:  py scripts/build_door_schedule_template.py [out.xlsx]
Prints the byte size, which content/news/door-hardware-schedule-guide.json and
content/downloads.json both quote as sizeBytes.
"""

import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

INK = "1A1A1A"
MUTED = "6B6B6B"
RULE = "D8D8D8"
HEADER_FILL = "F2F0ED"
BLANK_ROWS = 60
# Fixed so a rebuild with no content change gives the same metadata.
STAMP = datetime(2026, 10, 13)

thin = Side(style="thin", color=RULE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)

COLUMNS = [
    ("Door number", 11, "As on the drawings, e.g. D-101"),
    ("Location", 20, "Room or corridor; both sides if they differ"),
    ("Leaf / frame material", 18, "e.g. timber / steel, glass / aluminum"),
    ("Door size (W x H, mm)", 16, "Leaf width x height"),
    ("Thickness (mm)", 11, "Leaf thickness at the lock"),
    ("Fire rating", 14, "Rating and standard, or none"),
    ("Handing", 12, "LH / RH / LHR / RHR, or closing direction and the side it is seen from"),
    ("Function", 16, "e.g. passage, privacy, office, classroom, storeroom"),
    ("Lockset model", 18, "Lock case and lever or knob, or 'to suit'"),
    ("Cylinder", 18, "Profile, length each side (e.g. 30/30), keying group"),
    ("Closer", 16, "Model or power size, and arm type"),
    ("Hinges", 14, "Quantity and size per leaf"),
    ("Exit device", 16, "Model and type, if on an escape route"),
    ("Stop", 12, "Floor or wall stop"),
    ("Notes", 28, "Pairs, finish exceptions, changes in this revision"),
]

HEADER_BLOCK = [
    ("Project", "Name and address of the building"),
    ("Schedule revision", "Write it on every reissue; quotes refer to it"),
    ("Date", ""),
    ("Finish (whole project)", "Write exceptions in the Notes column"),
    ("Master key system", "Yes / no; attach the keying schedule if there is one"),
    ("Contact", "Who answers questions about this schedule"),
]

HANDING = ["LH", "RH", "LHR", "RHR", "Clockwise closing", "Anticlockwise closing"]


def build(out: Path) -> int:
    wb = Workbook()
    ws = wb.active
    ws.title = "Door schedule"
    wb.properties.creator = "Canton Hyland"
    wb.properties.created = STAMP
    wb.properties.modified = STAMP

    ws.cell(row=1, column=1, value="Door hardware schedule").font = Font(bold=True, size=14, color=INK)
    ws.cell(
        row=2,
        column=1,
        value="Canton Hyland Hardware & Building Material Co., Ltd.  ·  cantonlock.com/news/door-hardware-schedule-guide/",
    ).font = Font(size=9, color=MUTED)
    note = ws.cell(
        row=3,
        column=1,
        value="One row per door. Fill in what you know and leave the rest blank; a blank gets asked about, a guess gets ordered.",
    )
    note.font = Font(size=10, color=INK)

    row = 5
    for label, hint in HEADER_BLOCK:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, size=10, color=INK)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        for c in range(3, 7):
            ws.cell(row=row, column=c).border = box
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
        ws.cell(row=row, column=7, value=hint).font = Font(size=9, color=MUTED)
        row += 1

    header_row = row + 1
    for index, (label, width, help_text) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=index, value=label)
        cell.font = Font(bold=True, color=INK, size=10)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = box
        ws.column_dimensions[get_column_letter(index)].width = width
        hint = ws.cell(row=header_row + 1, column=index, value=help_text)
        hint.font = Font(size=8, color=MUTED)
        hint.alignment = Alignment(wrap_text=True, vertical="top")
        hint.border = box
    ws.row_dimensions[header_row].height = 30
    ws.row_dimensions[header_row + 1].height = 36

    first = header_row + 2
    last = first + BLANK_ROWS - 1
    for r in range(first, last + 1):
        for c in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = box
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # A suggestion list, not a lock: a schedule in another convention can still be typed in.
    handing_col = get_column_letter(1 + [c[0] for c in COLUMNS].index("Handing"))
    dv = DataValidation(type="list", formula1='"' + ",".join(HANDING) + '"', allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(f"{handing_col}{first}:{handing_col}{last}")

    ws.freeze_panes = ws.cell(row=first, column=2)
    ws.print_title_rows = f"{header_row}:{header_row + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out.stat().st_size


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("public/downloads/door-hardware-schedule-template.xlsx")
    size = build(target)
    # Self-check: the file reopens and carries all fifteen columns in order.
    from openpyxl import load_workbook

    sheet = load_workbook(target).active
    labels = [c.value for c in sheet[5 + len(HEADER_BLOCK) + 1]][: len(COLUMNS)]
    assert labels == [c[0] for c in COLUMNS], labels
    print(f"{target}  {size} bytes")
